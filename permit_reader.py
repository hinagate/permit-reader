"""Permit Reader - batch OCR for EMSD Lift Use Permits (CPU only).

Reads every image in the input folder, extracts Location ID and Address
(plus Lift No. and Date of Expiry), and writes one Excel file per run into
the output folder. The labels it looks for, and any extra fields, are set
in labels.ini next to the exe (created with the defaults on first run).

Usage:
    PermitReader.exe                       uses .\\input and .\\output next to the exe
    PermitReader.exe <input> [-o <output>] [--labels FILE] [--workers N] [--no-pause]
"""
from __future__ import annotations

import argparse
import configparser
import math
import multiprocessing as mp
import os
import re
import statistics
import sys
import time
import traceback
from concurrent.futures import ProcessPoolExecutor, as_completed
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path

import numpy as np
from PIL import Image, ImageOps, UnidentifiedImageError

APP_NAME = "Permit Reader"
VERSION = "1.1.0"
IMAGE_EXTS = {".jpg", ".jpeg", ".png", ".bmp", ".tif", ".tiff", ".webp"}

MAX_SIDE = 2000          # images are downscaled to this before OCR
DET_FAST = 1280          # text-detection size for the first pass
DET_FINE = 2000          # text-detection size for the retry pass
CPUS_PER_WORKER = 2      # many small workers beat one big one; ~300 MB RAM each
MAX_WORKERS = 8
MAX_THREADS = 4

LOC_MIN_SCORE = 0.85     # below this the Location ID is flagged for checking
ADDR_MIN_SCORE = 0.75


# --------------------------------------------------------------------------
# Image loading
# --------------------------------------------------------------------------

def load_image(path: Path) -> np.ndarray:
    """Load an image as a BGR array, honouring EXIF rotation, max side MAX_SIDE."""
    with Image.open(path) as im:
        w, h = im.size
        scale = MAX_SIDE / max(w, h)
        if scale < 1 and im.format == "JPEG":
            # Let the JPEG decoder skip detail we would throw away anyway.
            im.draft("RGB", (int(w * scale), int(h * scale)))
        im = ImageOps.exif_transpose(im).convert("RGB")
        if max(im.size) > MAX_SIDE:
            im.thumbnail((MAX_SIDE, MAX_SIDE), Image.Resampling.BILINEAR, reducing_gap=2.0)
        return np.ascontiguousarray(np.asarray(im)[:, :, ::-1])


def rotate(img: np.ndarray, degrees: int) -> np.ndarray:
    return np.ascontiguousarray(np.rot90(img, k=-degrees // 90)) if degrees else img


# --------------------------------------------------------------------------
# OCR boxes
# --------------------------------------------------------------------------

@dataclass
class Box:
    text: str
    score: float
    x0: float
    y0: float
    x1: float
    y1: float
    kind: str | None = None   # label kind, None for value text

    @property
    def cx(self) -> float:
        return (self.x0 + self.x1) / 2

    @property
    def cy(self) -> float:
        return (self.y0 + self.y1) / 2

    @property
    def h(self) -> float:
        return self.y1 - self.y0


def to_boxes(result) -> tuple[list[Box], bool]:
    """Convert a RapidOCR result into boxes in a deskewed coordinate frame.

    Also reports whether the text runs vertically (photo taken sideways).
    RapidOCR still reads such lines, but the page layout is then rotated.
    """
    if result.boxes is None or result.txts is None:
        return [], False
    pts = np.asarray(result.boxes, dtype=float)          # (n, 4, 2): tl, tr, br, bl
    top = pts[:, 1] - pts[:, 0]
    side = pts[:, 3] - pts[:, 0]
    top_len = np.hypot(top[:, 0], top[:, 1])
    side_len = np.hypot(side[:, 0], side[:, 1])
    wide = top_len > 1.5 * side_len
    sideways = int((side_len > 1.5 * top_len).sum()) > int(wide.sum())
    angle = float(np.median(np.arctan2(top[wide, 1], top[wide, 0]))) if wide.any() else 0.0
    c, s = math.cos(-angle), math.sin(-angle)
    pts = pts @ np.array([[c, -s], [s, c]]).T
    boxes = [
        Box(str(t).strip(), float(sc), p[:, 0].min(), p[:, 1].min(), p[:, 0].max(), p[:, 1].max())
        for p, t, sc in zip(pts, result.txts, result.scores)
        if str(t).strip()
    ]
    return boxes, sideways


# --------------------------------------------------------------------------
# Labels (labels.ini)
# --------------------------------------------------------------------------

LABELS_FILE = "labels.ini"
BUILTIN_FIELDS = {"location id": "loc", "lift no": "lift", "address": "addr", "date of expiry": "expiry"}

DEFAULT_LABELS_INI = """\
; Permit Reader - label settings
;
; Each field is found by the label printed beside it on the permit (for example
; 地址 / Address). The text to the right of the label is the field's value.
;
;   labels = the label words to look for. Separate them with |
;
; - Labels are matched at the start of a line. Spaces and upper/lower case do
;   not matter, and one OCR spelling mistake in a long English word is allowed.
; - Traditional and Simplified forms of common Chinese characters both match.
; - The "OCR Text" sheet in the Excel file shows how each label was read.
; - Save as UTF-8 (Notepad's default). Delete this file to restore the defaults.

; ---- Built-in fields -------------------------------------------------------
; You can change the label words, but not the [names].

[Location ID]
labels = 地點編號 | 地點 | Location ID | Location

[Lift No.]
labels = 升降機編號 | 自動梯編號 | Lift No | Escalator No

[Address]
labels = 地址 | Address

[Date of Expiry]
labels = 屆滿日期 | 屆滿 | Date of Expiry | Expiry

; ---- Your own fields -------------------------------------------------------
; Add one section per field. The [name] becomes the Excel column title.
;   type = text   the text to the right of the label, on the same line (default)
;   type = date   a Day / Month / Year date laid out like the expiry date
; A field that is not found is left blank and listed in the Note column.
;
; Example - remove the ; from the next three lines to add the date of issue:
;[Date of Issue]
;labels = 發出日期 | Date of Issue
;type = date
"""

# Traditional -> Simplified for characters common on permits. The OCR model
# mixes both forms, so labels are compared in Simplified.
SIMPLIFIED = str.maketrans(dict("""
    點点 編编 號号 機机 動动 屆届 滿满 發发 證证 準准 擁拥 責责 負负 載载 處处 業业 區区
    註注 冊册 員员 師师 廠厂 間间 層层 樓楼 電电 級级 數数 類类 檢检 驗验 簽签 應应 單单
    運运 維维 種种 條条 規规 範范 額额 輛辆 車车 門门 內内 屬属 產产 東东 認认 書书 記记
    錄录 項项 簡简 稱称 國国 際际 時时 開开 關关 標标 場场 館馆 廈厦 會会 總总 營营 經经
    為为 與与 於于 達达 過过 進进 統统 監监 測测 試试 備备 設设 裝装 終终 結结 據据 碼码
    齡龄 狀状 態态 歷历 線线 壓压 鋼钢 鐵铁 頂顶 戶户 灣湾 龍龙 華华 廣广 興兴 實实 園园
    製制 辦办 聯联 絡络 話话 傳传 續续
""".split()))
CJK = re.compile(r"[㐀-鿿豈-﫿]")


class SettingsError(Exception):
    pass


@dataclass
class LabelConfig:
    words: dict[str, tuple[str, ...]]   # label kind -> label words
    extra: dict[str, str]               # user field name -> "text" or "date"


def parse_labels(text: str, source: str, defaults: LabelConfig | None = None) -> LabelConfig:
    """Read label settings; built-in fields missing from the text keep their defaults."""
    cp = configparser.ConfigParser(interpolation=None, inline_comment_prefixes=(";",))
    try:
        cp.read_string(text, source)
    except configparser.Error as e:
        raise SettingsError(f"Problem in {source}: {e}") from None
    words = dict(defaults.words) if defaults else {}
    extra = {}
    for name in cp.sections():
        sec = cp[name]
        found = (w.strip().rstrip(".:：").strip() for w in sec.get("labels", "").split("|"))
        found = tuple(w for w in found if w)
        if not found:
            raise SettingsError(f"Problem in {source}: [{name}] needs label words, e.g. labels = 地址 | Address")
        kind = BUILTIN_FIELDS.get(name.strip().rstrip(".").lower())
        if kind:
            words[kind] = found
            continue
        ftype = sec.get("type", "text").strip().lower()
        if ftype not in ("text", "date"):
            raise SettingsError(f"Problem in {source}: [{name}] type must be text or date, not '{ftype}'")
        words[name] = found
        extra[name] = ftype
    return LabelConfig(words, extra)


DEFAULT_LABELS = parse_labels(DEFAULT_LABELS_INI, "default labels")


def load_labels(path: Path) -> LabelConfig:
    """Read the label settings file, creating it with the defaults if missing."""
    if not path.exists():
        try:
            # With a BOM, even old versions of Notepad show the Chinese correctly.
            path.write_text(DEFAULT_LABELS_INI, encoding="utf-8-sig")
        except OSError:   # read-only folder: run with the defaults
            return DEFAULT_LABELS
    try:
        text = path.read_text(encoding="utf-8-sig")
    except UnicodeDecodeError:
        raise SettingsError(f"{path.name} must be saved as UTF-8 "
                            "(in Notepad: File > Save As > Encoding: UTF-8)") from None
    return parse_labels(text, path.name, DEFAULT_LABELS)


def fold(text: str) -> tuple[str, list[int]]:
    """Lowercase, drop spaces and simplify Chinese; also map each kept char to its index."""
    out, index = [], []
    for i, ch in enumerate(text):
        if not ch.isspace():
            ch = ch.lower().translate(SIMPLIFIED)
            out.append(ch)
            index += [i] * len(ch)
    return "".join(out), index


def within_one_edit(a: str, b: str) -> bool:
    if abs(len(a) - len(b)) > 1:
        return False
    if len(a) > len(b):
        a, b = b, a
    i = 0
    while i < len(a) and a[i] == b[i]:
        i += 1
    return a[i + 1:] == b[i + 1:] if len(a) == len(b) else a[i:] == b[i + 1:]


_config = DEFAULT_LABELS
_matchers: list[tuple[str, str]] = []   # (kind, folded word), longest word first


def use_labels(config: LabelConfig) -> None:
    global _config, _matchers
    _config = config
    pairs = [(kind, fold(w)[0]) for kind, ws in config.words.items() for w in ws]
    _matchers = sorted(pairs, key=lambda p: -len(p[1]))


use_labels(DEFAULT_LABELS)


def match_label(text: str) -> tuple[str, int] | None:
    """Find a label at the start of an OCR line: (kind, end index in text).

    English labels must end at a word break ("No" does not match "Note");
    those of 6+ letters may have one OCR error.
    """
    t, index = fold(text)
    for kind, word in _matchers:
        fuzzy = word.isascii() and len(word) >= 6
        for n in (len(word), len(word) - 1, len(word) + 1) if fuzzy else (len(word),):
            if n > len(t) or not (within_one_edit(t[:n], word) if fuzzy else t[:n] == word):
                continue
            end = index[n - 1] + 1
            if not word.isascii() or not re.match(r"[a-z]", text[end:end + 1], re.I):
                return kind, end
    return None


def tag_labels(boxes: list[Box]) -> list[Box]:
    """Mark label boxes; split boxes where OCR merged a label with its value."""
    out = []
    for b in boxes:
        m = match_label(b.text)
        if not m:
            out.append(b)
            continue
        kind, end = m
        rest = b.text[end:].lstrip(" .:：")
        if rest:
            cut = b.x0 + (b.x1 - b.x0) * (len(b.text) - len(rest)) / len(b.text)
            out.append(Box(b.text[:end].strip(), b.score, b.x0, b.y0, cut, b.y1, kind))
            out.append(Box(rest, b.score, cut, b.y0, b.x1, b.y1))
        else:
            b.kind = kind
            out.append(b)
    return out


def label_pair(boxes: list[Box], kind: str, h: float) -> list[Box]:
    """First Chinese and first English label of a kind, if they sit together."""
    found = sorted((b for b in boxes if b.kind == kind), key=lambda b: b.cy)
    cn = next((b for b in found if CJK.search(b.text)), None)
    en = next((b for b in found if not CJK.search(b.text)), None)
    if cn and en and abs(cn.cy - en.cy) > 3 * h:
        return [en]
    return [b for b in (cn, en) if b]


def band(labels: list[Box], h: float, pad: float = 0.35) -> tuple[float, float]:
    return min(b.y0 for b in labels) - pad * h, max(b.y1 for b in labels) + pad * h


def group_lines(boxes: list[Box]) -> list[list[Box]]:
    lines: list[list[Box]] = []
    for b in sorted(boxes, key=lambda b: b.cy):
        if lines:
            ref = lines[-1]
            ref_cy = sum(x.cy for x in ref) / len(ref)
            if abs(b.cy - ref_cy) < 0.5 * max(b.h, max(x.h for x in ref)):
                ref.append(b)
                continue
        lines.append([b])
    return [sorted(line, key=lambda b: b.x0) for line in lines]


# --------------------------------------------------------------------------
# Field extraction
# --------------------------------------------------------------------------

DASH = r"[-‐‑–—−一_~=]"
LOC_STRICT = re.compile(rf"(?<!\d)(\d{{7}})\s*{DASH}\s*(\d{{3}})(?!\d)")
LOC_NODASH = re.compile(r"(?<!\d)(\d{7})(\d{3})(?!\d)")
LOC_LOOSE = re.compile(rf"(?<!\d)(\d{{4,9}})\s*{DASH}\s*(\d{{1,4}})(?!\d)")
DIGIT_LOOKALIKES = str.maketrans({"O": "0", "o": "0", "Q": "0", "I": "1", "l": "1", "|": "1",
                                  "i": "1", "!": "1", "S": "5", "s": "5", "B": "8", "Z": "2", "z": "2"})


@dataclass
class Fields:
    location_id: str = ""
    loc_strict: bool = False
    loc_score: float = 0.0
    address_cn: str = ""
    address_en: str = ""
    addr_score: float = 0.0
    lift_no: str = ""
    expiry: date | None = None
    extra: dict[str, str | date | None] = field(default_factory=dict)   # user fields
    anchors: int = 0          # how many built-in label kinds were found
    sideways: bool = False
    notes: tuple[str, ...] = ()
    ocr_text: str = ""

    @property
    def upright(self) -> bool:
        return self.anchors >= 2 and not self.sideways

    @property
    def complete(self) -> bool:
        return self.loc_strict and bool(self.address_cn) and bool(self.address_en)

    def rank(self) -> tuple:
        return (self.upright, self.loc_strict, bool(self.location_id),
                bool(self.address_cn) + bool(self.address_en),
                bool(self.lift_no) + (self.expiry is not None), self.anchors)


def extract_location(boxes, labels, h):
    """Return (location_id, strict_format, score, note, row_bottom)."""
    loc, lift = labels["loc"], labels["lift"]
    row, row_bottom = [], None
    if loc:
        top, row_bottom = band(loc, h)
        left = min(b.x1 for b in loc) - 0.3 * h
        right = min((b.x0 for b in lift if top <= b.cy <= row_bottom), default=math.inf)
        row = [b for b in boxes if not b.kind and top <= b.cy <= row_bottom
               and b.x0 >= left and b.cx < right]
        row.sort(key=lambda b: b.x0)

    def score_of(bs):
        return min((b.score for b in bs if re.search(r"\d", b.text)), default=0.0)

    raw = " ".join(b.text for b in row)
    for text, note in ((raw, ""), (raw.translate(DIGIT_LOOKALIKES), "Location ID: letters read as digits")):
        if m := LOC_STRICT.search(text):
            return f"{m[1]}-{m[2]}", True, score_of(row), note, row_bottom
    if m := LOC_NODASH.search(raw.replace(" ", "")):
        return f"{m[1]}-{m[2]}", True, score_of(row), "Location ID: dash not read, inserted", row_bottom
    # Label not found or row unreadable: look anywhere on the page.
    for b in sorted(boxes, key=lambda b: -b.score):
        if m := LOC_STRICT.search(b.text):
            return f"{m[1]}-{m[2]}", True, b.score, "Location ID found away from its label", row_bottom
    if m := LOC_LOOSE.search(raw):
        return f"{m[1]}-{m[2]}", False, score_of(row), "Location ID has an unusual format", row_bottom
    return "", False, 0.0, "", row_bottom


def extract_address(boxes, labels, h, loc_bottom):
    """Return (address_cn, address_en, min_score)."""
    addr, expiry = labels["addr"], labels["expiry"]
    if not addr:
        return "", "", 0.0
    top = min(b.y0 for b in addr) - 0.6 * h
    if loc_bottom is not None:
        top = max(top, loc_bottom)
    left = max(b.x1 for b in addr) - 0.3 * h
    bottom = min((b.y0 for b in expiry if b.y0 > top), default=None)

    cands = [b for b in boxes if not b.kind and b.cy > top and b.x0 >= left
             and (bottom is None or b.cy < bottom - 0.2 * h)]
    lines = group_lines(cands)
    if bottom is None:
        # No expiry label to stop at: keep lines until a large vertical gap.
        kept = lines[:1]
        for line in lines[1:6]:
            if line[0].cy - kept[-1][0].cy > 2.5 * h:
                break
            kept.append(line)
        lines = kept

    cn, en, used = [], [], []
    for line in lines:
        text = " ".join(b.text for b in line)
        if len(CJK.findall(text)) >= 2:
            cn.append("".join(b.text for b in line))
        else:
            en.append(text)
        used += line
    address_cn = re.sub(r"\s+", "", "".join(cn))
    address_en = " ".join(en).replace("，", ",")
    address_en = re.sub(r"\s+,", ",", re.sub(r"\s+", " ", address_en)).strip(" ,")
    return address_cn, address_en, min((b.score for b in used), default=0.0)


def extract_lift(boxes, labels, h):
    lift = labels["lift"]
    if not lift:
        return ""
    top, bot = band(lift, h)
    left = max(b.x1 for b in lift) - 0.3 * h
    vals = sorted((b for b in boxes if not b.kind and top <= b.cy <= bot and b.x0 >= left),
                  key=lambda b: b.x0)
    text = " ".join(b.text for b in vals).strip()
    return text.split()[0] if text else ""


def extract_text(boxes, found, h):
    """Text to the right of a label, on the label's line(s), up to the next label."""
    if not found:
        return ""
    top, bot = band(found, h)
    left = max(b.x1 for b in found) - 0.3 * h
    right = min((b.x0 for b in boxes if b.kind and b.x0 > left and top <= b.cy <= bot), default=math.inf)
    vals = [b for b in boxes if not b.kind and top <= b.cy <= bot and b.x0 >= left and b.cx < right]
    return " ".join(b.text for line in group_lines(vals) for b in line)


def extract_date(boxes, found, h):
    """A date read per column (Day / Month / Year) to avoid misreads.

    The numbers sit on the Chinese label's line, with the Day / Month / Year
    captions just below them. Another date block below (e.g. Date of Issue
    under Date of Expiry) has the same captions, so only the row directly
    under the numbers is used.
    """
    if not found:
        return None
    cn = [b for b in found if CJK.search(b.text)]
    if cn:
        top, bot = band(cn, h)
    else:   # only "Date of Expiry" found: numbers are on the line above it
        en = found[0]
        top, bot = en.y0 - 1.6 * en.h, en.y0 + 0.2 * en.h
    left = min(b.x1 for b in found) - 0.3 * h
    nums =[b for b in boxes if top <= b.cy <= bot and b.x0 >= left
            and re.fullmatch(r"\d{1,4}", b.text)]
    captions = [b for b in boxes if bot - 0.3 * h < b.cy < bot + 2 * h and b.x0 >= left]
    parts = {}
    for key, pat in (("d", r"day$"), ("m", r"month$"), ("y", r"year$")):
        col = min((b for b in captions if re.search(pat, b.text, re.I)), key=lambda b: b.cy, default=None)
        if col:
            best = min(nums, key=lambda b: abs(b.cx - col.cx), default=None)
            if best and abs(best.cx - col.cx) < (col.x1 - col.x0):
                parts[key] = int(best.text)
    if len(parts) < 3:
        tokens = re.findall(r"\d+", " ".join(b.text for b in sorted(nums, key=lambda b: b.x0)))
        if len(tokens) != 3:
            return None
        parts = dict(zip("dmy", map(int, tokens)))
    try:
        return date(parts["y"], parts["m"], parts["d"])
    except ValueError:
        return None


def parse(result) -> Fields:
    boxes, sideways = to_boxes(result)
    boxes = tag_labels(boxes)
    if not boxes:
        return Fields()
    h = statistics.median(b.h for b in boxes)
    labels = {k: label_pair(boxes, k, h) for k in _config.words}

    loc, strict, loc_score, note, loc_bottom = extract_location(boxes, labels, h)
    addr_cn, addr_en, addr_score = extract_address(boxes, labels, h, loc_bottom)
    return Fields(
        location_id=loc, loc_strict=strict, loc_score=loc_score,
        address_cn=addr_cn, address_en=addr_en, addr_score=addr_score,
        lift_no=extract_lift(boxes, labels, h),
        expiry=extract_date(boxes, labels["expiry"], h),
        extra={name: (extract_date if kind == "date" else extract_text)(boxes, labels[name], h)
               for name, kind in _config.extra.items()},
        anchors=sum(bool(labels[k]) for k in BUILTIN_FIELDS.values()),
        sideways=sideways,
        notes=(note,) if note else (),
        ocr_text="\n".join(" | ".join(b.text for b in line) for line in group_lines(boxes)),
    )


# --------------------------------------------------------------------------
# OCR engine and per-image work
# --------------------------------------------------------------------------

_engine = None


def disable_ort_spinning() -> None:
    """Stop onnxruntime threads busy-waiting between ops.

    Spinning threads fight each other when several OCR workers share the CPU;
    turning it off made batches 25-40% faster. RapidOCR has no setting for it,
    so wrap the method that builds its session options.
    """
    from rapidocr.inference_engine.onnxruntime.main import OrtInferSession
    build = OrtInferSession._init_sess_opts
    if getattr(build, "no_spin", False):
        return

    def no_spin_opts(cfg):
        opts = build(cfg)
        opts.add_session_config_entry("session.intra_op.allow_spinning", "0")
        return opts

    no_spin_opts.no_spin = True
    OrtInferSession._init_sess_opts = staticmethod(no_spin_opts)


def make_engine(threads: int):
    from rapidocr import RapidOCR
    disable_ort_spinning()
    return RapidOCR(params={
        "Global.log_level": "error",
        "Global.use_cls": False,          # whole-image rotation is handled in read_permit
        "Global.max_side_len": MAX_SIDE,
        "Det.limit_type": "max",
        "Det.limit_side_len": DET_FAST,
        "EngineConfig.onnxruntime.intra_op_num_threads": threads,
        "EngineConfig.onnxruntime.inter_op_num_threads": 1,
    })


def ocr(engine, img: np.ndarray, det_size: int) -> Fields:
    # RapidOCR rebuilds the detector's resize step from this on every call.
    engine._load_det_model().limit_side_len = det_size
    return parse(engine(img))


def read_permit(engine, img: np.ndarray) -> Fields:
    """OCR an image, retrying at other rotations or higher detail if needed."""
    best, best_img = ocr(engine, img, DET_FAST), img
    if not best.upright:
        if best.sideways:
            # Vertical text that still read correctly means the photo is turned
            # clockwise, so turning it back (270) is the likely fix.
            order = (270, 90) if best.anchors >= 2 else (90, 270)
        else:
            order = (180, 90, 270)
        for degrees in order:
            im = rotate(img, degrees)
            res = ocr(engine, im, DET_FAST)
            if res.rank() > best.rank():
                best, best_img = res, im
            if res.upright:
                break
    if not best.complete and max(best_img.shape[:2]) > DET_FAST:
        fine = ocr(engine, best_img, DET_FINE)
        if fine.rank() > best.rank():
            best = fine
    return best


def init_worker(threads: int, config: LabelConfig) -> None:
    global _engine
    use_labels(config)
    _engine = make_engine(threads)


def error_row(path_str: str, note: str = "") -> dict:
    return {"path": path_str, "location_id": "", "address_cn": "", "address_en": "",
            "lift_no": "", "expiry": None, "extra": {}, "status": "ERROR", "note": note,
            "ocr_text": "", "seconds": 0.0}


def process_file(path_str: str) -> dict:
    t0 = time.perf_counter()
    row = error_row(path_str)
    try:
        f = read_permit(_engine, load_image(Path(path_str)))
    except Exception as e:  # one bad file must not stop the batch
        row["note"] = ("Not a readable image file" if isinstance(e, UnidentifiedImageError)
                       else f"{type(e).__name__}: {e}")
        row["seconds"] = time.perf_counter() - t0
        return row

    notes = list(f.notes)
    if not f.location_id:
        notes.append("Location ID not found")
    elif f.loc_score < LOC_MIN_SCORE:
        notes.append(f"Location ID low confidence ({f.loc_score:.2f})")
    if not f.address_cn:
        notes.append("Chinese address not found")
    if not f.address_en:
        notes.append("English address not found")
    if (f.address_cn or f.address_en) and f.addr_score < ADDR_MIN_SCORE:
        notes.append(f"Address low confidence ({f.addr_score:.2f})")
    ok = f.loc_strict and f.address_cn and f.address_en and not any(
        n.startswith(("Location ID low", "Address low", "Location ID has")) for n in notes)
    notes += [f"{name} not found" for name, value in f.extra.items() if not value]

    row.update(location_id=f.location_id, address_cn=f.address_cn, address_en=f.address_en,
               lift_no=f.lift_no, expiry=f.expiry, extra=f.extra, status="OK" if ok else "CHECK",
               note="; ".join(notes), ocr_text=f.ocr_text, seconds=time.perf_counter() - t0)
    return row


# --------------------------------------------------------------------------
# Excel output
# --------------------------------------------------------------------------

def write_excel(rows: list[dict], out_path: Path, extra: dict[str, str]) -> None:
    from openpyxl import Workbook
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    def put(ws, r, c, value):
        if isinstance(value, str):
            value = ILLEGAL_CHARACTERS_RE.sub("", value)
        cell = ws.cell(row=r, column=c, value=value)
        if isinstance(value, str) and value.startswith("="):
            cell.data_type = "s"   # OCR text, never a formula
        return cell

    wb = Workbook()
    ws = wb.active
    ws.title = "Permits"
    columns = [("No.", 6, "center"), ("File", 36, None), ("Location ID", 15, "center"),
               ("Address (Chinese)", 42, "wrap"), ("Address (English)", 52, "wrap"),
               ("Lift No.", 9, "center"), ("Date of Expiry", 14, "center")]
    columns += [(name, 14, "center") if kind == "date" else (name, 30, "wrap")
                for name, kind in extra.items()]
    columns += [("Status", 9, "center"), ("Note", 45, "wrap")]
    head_fill = PatternFill("solid", fgColor="1F4E78")
    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    fills = {"CHECK": PatternFill("solid", fgColor="FFF2CC"), "ERROR": PatternFill("solid", fgColor="F8CBAD")}

    for c, (name, w, _) in enumerate(columns, 1):
        cell = put(ws, 1, c, name)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = head_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border
        ws.column_dimensions[get_column_letter(c)].width = w

    for i, r in enumerate(rows, 1):
        lift = int(r["lift_no"]) if re.fullmatch(r"[1-9]\d*", r["lift_no"]) else r["lift_no"]
        values = [i, r["file"], r["location_id"], r["address_cn"], r["address_en"],
                  lift, r["expiry"], *(r["extra"].get(name) for name in extra), r["status"], r["note"]]
        for c, (v, (_, _, align)) in enumerate(zip(values, columns), 1):
            cell = put(ws, i + 1, c, v)
            cell.border = border
            cell.alignment = Alignment(vertical="top", wrap_text=align == "wrap",
                                       horizontal="center" if align == "center" else None)
            if isinstance(v, date):
                cell.number_format = "dd/mm/yyyy"
            if r["status"] in fills:
                cell.fill = fills[r["status"]]
        if r.get("link"):
            link = ws.cell(row=i + 1, column=2)
            link.hyperlink = r["link"]
            link.font = Font(color="0563C1", underline="single")

    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{len(rows) + 1}"

    raw = wb.create_sheet("OCR Text")
    for c, (name, w) in enumerate((("File", 36), ("All text read from the image", 100)), 1):
        put(raw, 1, c, name).font = Font(bold=True)
        raw.column_dimensions[get_column_letter(c)].width = w
    for i, r in enumerate(rows, 2):
        put(raw, i, 1, r["file"]).alignment = Alignment(vertical="top")
        put(raw, i, 2, r["ocr_text"] or r["note"]).alignment = Alignment(vertical="top", wrap_text=True)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)


# --------------------------------------------------------------------------
# Batch runner
# --------------------------------------------------------------------------

def natural_key(p: Path):
    return [int(t) if t.isdigit() else t.lower() for t in re.split(r"(\d+)", str(p))]


def find_images(input_dir: Path, output_dir: Path) -> list[Path]:
    out = output_dir.resolve()
    files = [p for p in input_dir.rglob("*")
             if p.is_file() and p.suffix.lower() in IMAGE_EXTS and out not in p.resolve().parents]
    return sorted(files, key=lambda p: natural_key(p.relative_to(input_dir)))


def app_dir() -> Path:
    if getattr(sys, "frozen", False):
        return Path(sys.executable).resolve().parent
    return Path(__file__).resolve().parent


def run(input_dir: Path, output_dir: Path, workers: int | None, config: LabelConfig) -> int:
    files = find_images(input_dir, output_dir)
    print(f"Input : {input_dir}")
    print(f"Output: {output_dir}")
    if not files:
        print(f"\nNo images found. Put photos ({', '.join(sorted(IMAGE_EXTS))}) in the input folder.")
        return 1

    cpu = os.process_cpu_count() or 4
    if workers is None:
        workers = max(1, min(cpu // CPUS_PER_WORKER, MAX_WORKERS))
    workers = max(1, min(workers, len(files)))
    threads = max(1, min(MAX_THREADS, cpu // workers))
    print(f"Images: {len(files)}   (CPU: {workers} worker(s) x {threads} thread(s))\n")

    t0 = time.perf_counter()
    rows: dict[str, dict] = {}

    def report(row: dict) -> None:
        rows[row["path"]] = row
        rel = Path(row["path"]).relative_to(input_dir)
        loc = row["location_id"] or "-"
        print(f"[{len(rows):>4}/{len(files)}] {row['status']:<5} {loc:<12} {rel}  ({row['seconds']:.1f}s)")

    if workers == 1:
        init_worker(threads, config)
        for p in files:
            report(process_file(str(p)))
    else:
        with ProcessPoolExecutor(max_workers=workers, initializer=init_worker,
                                 initargs=(threads, config)) as pool:
            futures = {pool.submit(process_file, str(p)): str(p) for p in files}
            for fut in as_completed(futures):
                try:
                    report(fut.result())
                except Exception as e:   # worker process died (e.g. out of memory)
                    report(error_row(futures[fut], f"OCR worker failed: {type(e).__name__}"))

    ordered = []
    for p in files:
        row = rows[str(p)]
        row["file"] = str(p.relative_to(input_dir))
        try:
            row["link"] = os.path.relpath(p, output_dir)
        except ValueError:   # different drive
            row["link"] = str(p)
        ordered.append(row)

    out_path = output_dir / f"Permit_Results_{datetime.now():%Y%m%d_%H%M%S}.xlsx"
    write_excel(ordered, out_path, config.extra)

    elapsed = time.perf_counter() - t0
    counts = {s: sum(r["status"] == s for r in ordered) for s in ("OK", "CHECK", "ERROR")}
    print(f"\nDone: {len(files)} image(s) in {elapsed:.1f}s ({elapsed / len(files):.2f}s per image)")
    print(f"OK: {counts['OK']}   CHECK: {counts['CHECK']}   ERROR: {counts['ERROR']}")
    print(f"Saved: {out_path}")
    return 0


def main() -> int:
    for stream in (sys.stdout, sys.stderr):
        if stream and hasattr(stream, "reconfigure"):
            stream.reconfigure(encoding="utf-8", errors="replace")

    base = app_dir()
    ap = argparse.ArgumentParser(prog="PermitReader", description=f"{APP_NAME} {VERSION}")
    ap.add_argument("input", nargs="?", type=Path, default=base / "input",
                    help="folder with permit photos (default: input next to the program)")
    ap.add_argument("-o", "--output", type=Path, default=None,
                    help="folder for the Excel file (default: output next to the program)")
    ap.add_argument("--labels", type=Path, default=None,
                    help=f"label settings file (default: {LABELS_FILE} next to the program)")
    ap.add_argument("--workers", type=int, default=None, help="parallel OCR processes (default: auto)")
    ap.add_argument("--no-pause", action="store_true", help="do not wait for Enter at the end")
    args = ap.parse_args()

    input_dir = args.input.resolve()
    output_dir = (args.output or base / "output").resolve()
    labels_path = (args.labels or base / LABELS_FILE).resolve()
    print(f"{APP_NAME} {VERSION}\n")
    try:
        config = load_labels(labels_path)
        extra = f"  (extra fields: {', '.join(config.extra)})" if config.extra else ""
        print(f"Labels: {labels_path}{extra}")
        if not input_dir.is_dir():
            input_dir.mkdir(parents=True, exist_ok=True)
            print(f"Created input folder: {input_dir}\nPut permit photos in it and run again.")
            code = 1
        else:
            code = run(input_dir, output_dir, args.workers, config)
    except SettingsError as e:
        print(e)
        code = 2
    except KeyboardInterrupt:
        print("\nCancelled.")
        code = 130
    except Exception:
        traceback.print_exc()
        code = 2

    if not args.no_pause and sys.stdin and sys.stdin.isatty():
        try:
            input("\nPress Enter to close...")
        except EOFError:
            pass
    return code


if __name__ == "__main__":
    mp.freeze_support()
    sys.exit(main())
